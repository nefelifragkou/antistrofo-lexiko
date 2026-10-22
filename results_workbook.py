"""Excel με τα αποτελέσματα της αξιολόγησης (λεξικό vs LLM) — βήματα 6–7."""
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

root = Path(__file__).resolve().parent
rows = json.load(open(root / "comparison.json", encoding="utf-8"))
import csv
D = {int(d["id"]): d for d in csv.DictReader(open(root / "eval_results_dictionary.csv", encoding="utf-8-sig"))}
out = sys.argv[1]
F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
def f(**k): return Font(name=F, **k)

wb = Workbook()
ws = wb.active; ws.title = "Σύνοψη"
dt = wb.create_sheet("Ανά περιγραφή")
hdr = ["#", "Λέξη-στόχος", "Περιγραφή", "Είδος", "Δυσκολία", "Ίδια ρίζα", "Θέση — λεξικό", "Θέση — ChatGPT",
       "Θέση — Gemini", "10 πρώτες — λεξικό", "10 πρώτες — ChatGPT", "10 πρώτες — Gemini",
       "ChatGPT: λέξεις που υπάρχουν στο λεξικό (από 10)", "Gemini: λέξεις που υπάρχουν στο λεξικό (από 10)"]
widths = [5, 16, 46, 12, 10, 9, 10, 10, 10, 44, 44, 44, 14, 14]
for j, (h, w) in enumerate(zip(hdr, widths), 1):
    c = dt.cell(1, j, h); c.font = H; c.fill = HF; c.border = B; c.alignment = Alignment(wrap_text=True, vertical="center")
    dt.column_dimensions[chr(64 + j)].width = w
for i, r in enumerate(rows, 2):
    d = D[r["id"]]
    vals = [r["id"], r["target"], d["description"], r["pos"], r["difficulty"], r["same_root"],
            r["dict_rank"], r["llm_rank"], r["gem_rank"], d["top10"], r["llm_top10"], r["gem_top10"],
            r["llm_valid"], r["gem_valid"]]
    for j, v in enumerate(vals, 1):
        c = dt.cell(i, j, v); c.font = f(); c.border = B; c.alignment = wrap
    for j in (7, 8, 9):
        v = dt.cell(i, j).value
        dt.cell(i, j).fill = PatternFill("solid", fgColor="E5F4E4" if v == 1 else ("FFF2CC" if v else "FBEBEB"))
n = len(rows) + 1
dt.freeze_panes = "C2"
dt.auto_filter.ref = f"A1:N{n}"

ws.column_dimensions["A"].width = 40
for col in "BCDEFGHI": ws.column_dimensions[col].width = 14
ws["A1"] = "Αξιολόγηση — Αντίστροφο Λεξικό έναντι LLM"; ws["A1"].font = f(bold=True, size=14)
ws["A2"] = "30 περιγραφές ελέγχου. Θέση = σε ποια θέση της δεκάδας βγήκε η σωστή λέξη (κενό = δεν βγήκε). Οι τιμές υπολογίζονται από το φύλλο «Ανά περιγραφή»."
ws["A2"].font = f(italic=True, color="595959")
ws["A3"] = "LLM: ChatGPT και Gemini (δωρεάν, χωρίς σύνδεση), 22/10/2026 — χωρίς πρόσβαση στο λεξικό, ίδιο prompt. Λεξικό: intfloat/multilingual-e5-large (ONNX)."
ws["A3"].font = f(italic=True, color="595959")
heads = ["Ομάδα", "Πλήθος", "Λεξικό: σωστή 1η", "Λεξικό: στις 10", "ChatGPT: σωστή 1η", "ChatGPT: στις 10",
         "Gemini: σωστή 1η", "Gemini: στις 10"]
for j, h in enumerate(heads, 1):
    c = ws.cell(5, j, h); c.font = H; c.fill = HF; c.border = B; c.alignment = Alignment(wrap_text=True, vertical="center")
R = "'Ανά περιγραφή'!"
groups = [("Σύνολο", None, None), ("Χωρίς τις 4 «ίδιας ρίζας»", "F", "όχι"),
          ("Εύκολες", "E", "Εύκολη"), ("Μέτριες", "E", "Μέτρια"), ("Δύσκολες", "E", "Δύσκολη"),
          ("Ουσιαστικά", "D", "Ουσιαστικό"), ("Επίθετα", "D", "Επίθετο"), ("Ρήματα", "D", "Ρήμα")]
for i, (name, col, val) in enumerate(groups, 6):
    ws.cell(i, 1, name)
    if col:
        crit = f'{R}${col}$2:${col}${n},"{val}"'
        ws.cell(i, 2, f"=COUNTIF({crit})")
        ws.cell(i, 3, f"=COUNTIFS({crit},{R}$G$2:$G${n},1)/B{i}")
        ws.cell(i, 4, f'=COUNTIFS({crit},{R}$G$2:$G${n},"<>")/B{i}')
        ws.cell(i, 5, f"=COUNTIFS({crit},{R}$H$2:$H${n},1)/B{i}")
        ws.cell(i, 6, f'=COUNTIFS({crit},{R}$H$2:$H${n},"<>")/B{i}')
        ws.cell(i, 7, f"=COUNTIFS({crit},{R}$I$2:$I${n},1)/B{i}")
        ws.cell(i, 8, f'=COUNTIFS({crit},{R}$I$2:$I${n},"<>")/B{i}')
    else:
        ws.cell(i, 2, f"=COUNTA({R}$B$2:$B${n})")
        ws.cell(i, 3, f"=COUNTIF({R}$G$2:$G${n},1)/B{i}")
        ws.cell(i, 4, f"=COUNT({R}$G$2:$G${n})/B{i}")
        ws.cell(i, 5, f"=COUNTIF({R}$H$2:$H${n},1)/B{i}")
        ws.cell(i, 6, f"=COUNT({R}$H$2:$H${n})/B{i}")
        ws.cell(i, 7, f"=COUNTIF({R}$I$2:$I${n},1)/B{i}")
        ws.cell(i, 8, f"=COUNT({R}$I$2:$I${n})/B{i}")
    for j in range(1, 9):
        c = ws.cell(i, j); c.border = B; c.font = f(bold=(i == 6))
        if j >= 3: c.number_format = "0%"
ws.cell(15, 1, "Λέξεις που υπάρχουν στο λεξικό — ChatGPT"); ws.cell(15, 1).font = f(bold=True)
ws.cell(15, 2, f"=SUM({R}$M$2:$M${n})/(10*B6)"); ws.cell(15, 2).number_format = "0%"
ws.cell(17, 1, "Λέξεις που υπάρχουν στο λεξικό — Gemini"); ws.cell(17, 1).font = f(bold=True)
ws.cell(17, 2, f"=SUM({R}$N$2:$N${n})/(10*B6)"); ws.cell(17, 2).number_format = "0%"
ws.cell(16, 1, "Οι υπόλοιπες είναι κυρίως μετοχές, θηλυκοί τύποι ή φράσεις, αλλά και λέξεις που δεν υπάρχουν (π.χ. «ταχυδρόμης», «αιμοδοτούχος», «αυτοδιδάκτωρ», «καθρεπτίζουσα») ή ξένες («pastrychef»).")
ws.cell(16, 1).font = f(italic=True, color="595959")
wb.save(out)
print("ok")
