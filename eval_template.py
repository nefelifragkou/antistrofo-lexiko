"""Φτιάχνει το Excel για τις 30 περιγραφές ελέγχου (βήμα 6)."""
import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

root = Path(__file__).resolve().parent
S = [json.loads(l) for l in open(root / "data/processed/senses.jsonl", encoding="utf-8")]
dev = json.load(open(root / "dev_queries.json", encoding="utf-8"))
words = sorted({s["headword"] for s in S} | {a for s in S for a in s["alt_forms"]})
out = sys.argv[1]

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")
Y = PatternFill("solid", fgColor="FFF2CC"); G = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
def f(**k): return Font(name=F, **k)

wb = Workbook()
ws = wb.active; ws.title = "Οδηγίες"
ws.column_dimensions["A"].width = 105
lines = [
    ("30 περιγραφές ελέγχου — Αντίστροφο Λεξικό", f(bold=True, size=14)),
    ("", None),
    ("Τι κάνεις", f(bold=True, size=12)),
    ("1. Στο φύλλο «Περιγραφές» γράψε 30 λέξεις (στήλη B) και για καθεμία μια περιγραφή με δικά σου λόγια (στήλη C).", f()),
    ("2. Διάλεξε είδος λέξης (D) και δυσκολία (E) από τις λίστες.", f()),
    ("3. Η στήλη F ελέγχει μόνη της ότι η λέξη υπάρχει στο λεξικό. Αν γράφει ✘, έλεγξε τόνο και γραφή (π.χ. «ορφανός», όχι «ορφανό»).", f()),
    ("", None),
    ("Κανόνες για να είναι δίκαιο το τεστ", f(bold=True, size=12)),
    ("• Γράψε την περιγραφή ΠΡΙΝ δεις τον ορισμό του λεξικού — όπως θα την έγραφε κάποιος που ψάχνει τη λέξη.", f()),
    ("• Μη δοκιμάσεις τις περιγραφές στη σελίδα πριν τις «κλειδώσουμε».", f()),
    ("• Μη χρησιμοποιήσεις μέσα στην περιγραφή την ίδια τη λέξη ή λέξη της ίδιας ρίζας (π.χ. όχι «φοβία» για την κλειστοφοβία).", f()),
    ("• Μη χρησιμοποιήσεις τις 20 λέξεις που δοκιμάσαμε ήδη (η στήλη F θα σε ειδοποιήσει).", f()),
    ("", None),
    ("Καλή ανάμειξη (στόχος)", f(bold=True, size=12)),
    ("• Δυσκολία: περίπου 10 εύκολες (καθημερινές λέξεις), 12 μέτριες, 8 δύσκολες (σπάνιες ή λόγιες).", f()),
    ("• Είδος: κυρίως ουσιαστικά, αλλά και αρκετά επίθετα και ρήματα (π.χ. 18 / 6 / 6).", f()),
    ("• Ποικιλία θεμάτων: αντικείμενα, συναισθήματα, επαγγέλματα, φύση, αφηρημένες έννοιες.", f()),
    ("", None),
    ("Υπόμνημα: κίτρινο = συμπληρώνεις εσύ · γκρι = παράδειγμα · λευκό = υπολογίζεται μόνο του", f(italic=True, color="595959")),
]
for i, (t, ft) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.alignment = Alignment(wrap_text=True, vertical="top")
    if ft: c.font = ft

# Λίστες αναφοράς
wl = wb.create_sheet("Λήμματα")
wl["A1"] = "Λέξεις του λεξικού (για τον έλεγχο της στήλης F — μην αλλάξεις)"; wl["A1"].font = f(bold=True)
for i, w in enumerate(words, 2):
    wl.cell(i, 1, w)
wl["C1"] = "Λέξεις που χρησιμοποιήθηκαν ήδη στη δοκιμή μοντέλων"; wl["C1"].font = f(bold=True)
for i, (_, w) in enumerate(dev, 2):
    wl.cell(i, 3, w)
wl.column_dimensions["A"].width = 30; wl.column_dimensions["C"].width = 30
last_w, last_d = len(words) + 1, len(dev) + 1

ws = wb.create_sheet("Περιγραφές", 1)
hdr = ["#", "Λέξη-στόχος", "Περιγραφή (με δικά μου λόγια)", "Είδος λέξης", "Δυσκολία", "Έλεγχος λέξης", "Σχόλιο"]
widths = [5, 20, 60, 14, 12, 34, 30]
for j, (h, w) in enumerate(zip(hdr, widths), 1):
    c = ws.cell(3, j, h); c.font = H; c.fill = HF; c.border = B; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[chr(64 + j)].width = w
ws["A1"] = "Συμπληρωμένες:"; ws["A1"].font = f(bold=True)
ws["C1"] = '=COUNTIFS(B5:B34,"<>",C5:C34,"<>")&" από 30   ·   εύκολες: "&COUNTIF(E5:E34,"Εύκολη")&"   μέτριες: "&COUNTIF(E5:E34,"Μέτρια")&"   δύσκολες: "&COUNTIF(E5:E34,"Δύσκολη")&"   ·   ουσ.: "&COUNTIF(D5:D34,"Ουσιαστικό")&"   επίθ.: "&COUNTIF(D5:D34,"Επίθετο")&"   ρήματα: "&COUNTIF(D5:D34,"Ρήμα")'
ws["C1"].font = f(bold=True, color="1F4E78")
ws.merge_cells("C1:G1")

ex = ["π.χ.", "ομπρέλα", "αντικείμενο που ανοίγουμε πάνω από το κεφάλι για να μη βραχούμε", "Ουσιαστικό", "Εύκολη", None, "Παράδειγμα — μην το χρησιμοποιήσεις"]
for j, v in enumerate(ex, 1):
    c = ws.cell(4, j, v); c.fill = G; c.border = B; c.alignment = wrap; c.font = f(italic=True, color="595959")
dv_pos = DataValidation(type="list", formula1='"Ουσιαστικό,Επίθετο,Ρήμα,Επίρρημα"', allow_blank=True)
dv_dif = DataValidation(type="list", formula1='"Εύκολη,Μέτρια,Δύσκολη"', allow_blank=True)
ws.add_data_validation(dv_pos); ws.add_data_validation(dv_dif)
for r in range(4, 35):
    if r > 4:
        ws.cell(r, 1, r - 4)
        for j in (2, 3, 4, 5, 7):
            c = ws.cell(r, j); c.fill = Y; c.border = B; c.alignment = wrap; c.font = f()
        ws.cell(r, 1).font = f(); ws.cell(r, 1).border = B
    ws.cell(r, 6, f'=IF(B{r}="","",IF(COUNTIF(Λήμματα!$C$2:$C${last_d},B{r})>0,"⚠ χρησιμοποιήθηκε ήδη στη δοκιμή",'
                  f'IF(COUNTIF(Λήμματα!$A$2:$A${last_w},B{r})>0,"✔ υπάρχει στο λεξικό","✘ δεν βρέθηκε — έλεγξε τόνο/γραφή")))')
    ws.cell(r, 6).border = B; ws.cell(r, 6).font = f(); ws.cell(r, 6).alignment = wrap
    ws.row_dimensions[r].height = 32
    dv_pos.add(f"D{r}"); dv_dif.add(f"E{r}")
red = PatternFill("solid", fgColor="FBEBEB"); green = PatternFill("solid", fgColor="E5F4E4")
ws.conditional_formatting.add("F5:F34", FormulaRule(formula=['LEFT(F5,1)="✔"'], fill=green))
ws.conditional_formatting.add("F5:F34", FormulaRule(formula=['OR(LEFT(F5,1)="✘",LEFT(F5,1)="⚠")'], fill=red))
ws.freeze_panes = "A5"
wb.active = 1
wb.save(out)
print("ok", len(words))
