"""Φτιάχνει το Excel ελέγχου του βήματος 3 (σημασίες των 31 λημμάτων + τυχαίο δείγμα)."""
import json, random, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

root = Path(__file__).resolve().parent.parent
S = [json.loads(l) for l in open(root / "data/processed/senses.jsonl", encoding="utf-8")]
E = [json.loads(l) for l in open(root / "data/processed/excluded.jsonl", encoding="utf-8")]
stats = json.load(open(root / "data/processed/stats.json", encoding="utf-8"))
test_lemmas = json.load(open(sys.argv[1], encoding="utf-8"))  # λίστα λημμάτων (μετά τη διόρθωση γραμμάτων)
out = sys.argv[2]

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")
Y = PatternFill("solid", fgColor="FFF2CC"); G = PatternFill("solid", fgColor="EDEDED")
thin = Side(style="thin", color="BFBFBF"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

def f(**k): return Font(name=F, **k)

wb = Workbook()
ws = wb.active; ws.title = "Σύνοψη"
ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 70
ws["A1"] = "Βήμα 3 — Χωρισμός λημμάτων σε σημασίες"; ws["A1"].font = f(bold=True, size=14)
ws["A2"] = "Αποτέλεσμα του κώδικα src/split_senses.py σε όλο το λεξικό."; ws["A2"].font = f(italic=True, color="595959")
rows = [("Μέτρηση", "Τιμή", "Σχόλιο"),
        ("Λήμματα στο λεξικό", stats["entries"], ""),
        ("Λήμματα που μπήκαν στο αντίστροφο λεξικό", stats["lemmas_with_senses"], ""),
        ("Σημασίες (γραμμές για αναζήτηση)", stats["senses"], "Αυτές θα γίνουν «αριθμοί νοήματος» (embeddings) στο βήμα 4"),
        ("   από τις οποίες ειδικότερες χρήσεις μετά το ||", stats["sub_uses"], "π.χ. χορτοφάγος: «(για άνθρ.) που συνηθίζει να τρώει χόρτα…»"),
        ("Δανεικοί ορισμοί", stats["borrowed_definitions"], "Ορισμός μίας λέξης που πήρε τον ορισμό της λέξης στην οποία παραπέμπει (απόφαση 3)"),
        ("Εξαιρέθηκαν: προθήματα / επιθήματα", stats["excluded"].get("πρόθημα/επίθημα", 0), "Απόφαση 2"),
        ("Εξαιρέθηκαν: λέξεις γραμματικής / επιφωνήματα", stats["excluded"].get("λέξη γραμματικής/επιφώνημα", 0), "Απόφαση 2 — η λίστα στο φύλλο «Εξαιρέθηκαν»"),
        ("Εξαιρέθηκαν: χωρίς ορισμό μετά τον χωρισμό", stats["excluded"].get("χωρίς ορισμό μετά τον χωρισμό", 0), "Λήμματα με μόνο παραδείγματα ή φράσεις"),
        ]
for i, r in enumerate(rows, 4):
    for j, v in enumerate(r, 1):
        c = ws.cell(i, j, v); c.border = B; c.alignment = wrap
        c.font = H if i == 4 else f()
        if i == 4: c.fill = HF
        if isinstance(v, int): c.number_format = "#,##0"
ws.cell(15, 1, "Τι να ελέγξεις").font = f(bold=True, size=12)
notes = ["1. Φύλλο «31 λήμματα»: οι σημασίες που έβγαλε ο υπολογιστής για τα λήμματα που εξερεύνησες. Σύγκρινε με τις σημειώσεις σου.",
         "2. Φύλλο «Τυχαίο δείγμα»: 50 τυχαίες σημασίες από όλο το λεξικό — δείχνουν αν οι κανόνες δουλεύουν και σε λήμματα που δεν είδαμε.",
         "3. Στις κίτρινες στήλες: «Σωστό;» (Ναι / Όχι / Μισό) και σχόλιο αν κάτι λείπει ή περισσεύει.",
         "Κωδικός σημασίας: I/II = ομάδα, 1/2 = σημασία, α/β = υποσημασία, + = ειδικότερη χρήση μετά το ||, 0 = γενικός ορισμός πριν την αρίθμηση."]
for i, n in enumerate(notes, 16):
    c = ws.cell(i, 1, n); c.font = f(); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3); ws.row_dimensions[i].height = 32

hdr = ["Λήμμα", "Κωδικός", "Ετικέτες", "Ορισμός (όπως θα αναζητείται)", "Συνώνυμα", "Αντώνυμα", "Δανεικός από", "Σωστό;", "Σχόλιο"]
widths = [22, 9, 20, 80, 18, 16, 14, 10, 30]

def sense_sheet(title, senses, example_first=False):
    ws = wb.create_sheet(title)
    for j, (h, w) in enumerate(zip(hdr, widths), 1):
        c = ws.cell(1, j, h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = B
        ws.column_dimensions[chr(64 + j)].width = w
    dv = DataValidation(type="list", formula1='"Ναι,Όχι,Μισό"', allow_blank=True); ws.add_data_validation(dv)
    prev = None
    for i, s in enumerate(senses, 2):
        new = s["lemma"] != prev or s["code"] in ("0", "1", "1α", "I", "I1", "I1α")
        vals = [s["lemma"] if new else "", s["code"], ", ".join(s["labels"]), s["definition"],
                ", ".join(s["synonyms"]), ", ".join(s["antonyms"]), s.get("borrowed_from", "")]
        prev = s["lemma"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = f(bold=(j == 1)); c.alignment = wrap; c.border = B
        for j in (8, 9):
            c = ws.cell(i, j); c.fill = Y; c.border = B; c.alignment = wrap; c.font = f()
        dv.add(f"H{i}")
    if example_first and senses:
        ws.cell(2, 8, "Ναι").fill = G; ws.cell(2, 9, "Παράδειγμα: σωστός ορισμός, χωρίς το αντώνυμο μέσα").fill = G
        ws.cell(2, 8).font = f(italic=True, color="595959"); ws.cell(2, 9).font = f(italic=True, color="595959")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:I{len(senses) + 1}"

order = {l: i for i, l in enumerate(test_lemmas)}
t31 = sorted([s for s in S if s["lemma"] in order], key=lambda s: (order[s["lemma"]], s["id"]))
sense_sheet("31 λήμματα", t31, example_first=True)
random.seed(2026)
sample = sorted(random.sample(S, 50), key=lambda s: s["id"])
sense_sheet("Τυχαίο δείγμα", sample)

ws = wb.create_sheet("Εξαιρέθηκαν")
for j, (h, w) in enumerate(zip(["Λήμμα", "Λόγος"], [30, 40]), 1):
    c = ws.cell(1, j, h); c.font = H; c.fill = HF; c.border = B
    ws.column_dimensions[chr(64 + j)].width = w
ex = [e for e in E if e["reason"] != "πρόθημα/επίθημα"]
for i, e in enumerate(ex, 2):
    ws.cell(i, 1, e["lemma"]).font = f(); ws.cell(i, 2, e["reason"]).font = f()
ws.cell(len(ex) + 3, 1, f"Τα {stats['excluded'].get('πρόθημα/επίθημα', 0):,} προθήματα/επιθήματα δεν εμφανίζονται εδώ (όλα τα λήμματα που αρχίζουν ή τελειώνουν με «-»).".replace(",", ".")).font = f(italic=True, color="595959")
ws.freeze_panes = "A2"
wb.save(out)
print("ok", len(t31), len(sample), len(ex))
